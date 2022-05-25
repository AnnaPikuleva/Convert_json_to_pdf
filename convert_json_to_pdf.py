import argparse
parser = argparse.ArgumentParser(description='Json для преобразования в pdf')
parser.add_argument('indir', type=str, help='Путь к таблице')
parser.add_argument('outdir', type=str, help='Путь к папке где сохранить файл')
args = parser.parse_args()
print(args.indir)

import json
import pandas as pd
import docx
from docx2pdf import convert

js = r"args.indir"
df = pd.read_json(js)

doc = docx.Document() 
table = doc.add_table(rows=df.shape[0], cols=df.shape[1])
for i in range(df.shape[0]):
    for j in range(df.shape[1]):
        table.cell(i,j).text = str(df.values[i,j])
doc = doc.save("output_file_path.docx")

path = "output_file_path.docx"
convert(path, r"args.outdir")
